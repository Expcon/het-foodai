"""生成《真实八类蔬菜模型 Gradio Demo 端到端验收交付文档》Markdown + DOCX。"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
EVID = ROOT / "docs" / "acceptance-evidence" / "20261021"
OUT = ROOT.parent / "deliverables"
OUT.mkdir(parents=True, exist_ok=True)

cpu = json.loads((EVID / "acceptance.json").read_text(encoding="utf-8"))
gpu = json.loads((EVID / "acceptance_gpu.json").read_text(encoding="utf-8"))
classes = json.loads((ROOT / "configs" / "classes.json").read_text(encoding="utf-8"))["classes"]
cpu_by = {r["file"]: r for r in cpu["images"]}
gpu_by = {r["file"]: r for r in gpu["images"]}

def rng(records, key):
    vals = [r[key] for r in records]
    return f"{min(vals):.1f} – {max(vals):.1f}"

TITLE = "真实八类蔬菜模型 Gradio Demo 端到端验收交付文档"
SUBTITLE = "HeT FoodAI · Dataset v1.1.0 · CPU / GPU 实测 · 2026-10-21"

# ------------------------------------------------------------------ Markdown
L = []
L.append(f"# {TITLE}")
L.append("")
L.append(f"*{SUBTITLE}*")
L.append("")
L.append("本文档记录真实 EfficientNetV2-S 八类蔬菜模型接入 Gradio Demo 的端到端验收结果，"
         "按交付要求分为五项：运行截图、真实图片预测结果、Top-3 与时延及异常处理记录、"
         "代码问题与修复、课程展示结论。所有数字均为本机实测，softmax 分数不等同于模型准确率。")
L.append("")
L.append("## 0. 验收对象与环境")
L.append("")
L.append("| 项目 | 内容 |")
L.append("|---|---|")
L.append(f"| 模型 | `{cpu['model_id']}`（EfficientNetV2-S，stage 8） |")
L.append(f"| 输入尺寸 | {cpu['input_size']} px |")
L.append(f"| 冻结类别配置 | `configs/classes.json`，status = frozen，version = v1.1.0 |")
L.append(f"| class_list_sha256 | `{cpu['class_list_sha256']}` |")
L.append(f"| checkpoint SHA256 | `{cpu['checkpoint_sha256']}` |")
L.append(f"| CPU 环境 | Windows · Python {cpu['python']} · torch {cpu['torch']} · device=cpu |")
L.append(f"| GPU 环境 | Windows · Python {gpu['python']} · torch {gpu['torch']} · "
         f"device=cuda（NVIDIA GeForce RTX 4050 Laptop GPU） |")
L.append("")
L.append("## 1. Gradio 正常运行截图")
L.append("")
L.append("| 截图 | 说明 |")
L.append("|---|---|")
L.append("| `demo_ready_cucumber.png` | CPU 环境，真实黄瓜图片识别出 Top-3，界面显示 110.3 ms |")
L.append("| `demo_ready_gpu_tomato.png` | GPU 环境（设备 cuda），真实西红柿图片识别出 Top-3 |")
L.append("| `demo_corrupt.png` | 损坏图片：明确错误提示，无任何预测分数 |")
L.append("| `demo_no_upload.png` | 未上传图片点击识别：提示“请先上传图片” |")
L.append("| `demo_missing_checkpoint.png` | 未配置模型：显示未就绪，识别按钮禁用 |")
L.append("")
L.append("## 2. 真实图片预测结果")
L.append("")
L.append("8 张独立真实图片，CPU 与 GPU 预测结果一致；每张重复 5 次（CPU）/ 10 次（GPU），"
         "时延为中位数，单位 ms。")
L.append("")
L.append("| 图片 | Top-1 | Top-2 | Top-3 | CPU 时延 | GPU 时延 |")
L.append("|---|---|---|---|---:|---:|")
for f, c in cpu_by.items():
    g = gpu_by[f]
    ct, gt = c["top3"], g["top3"]
    L.append(
        f"| {f} | {ct[0]['label_zh']} ({ct[0]['class_key']}) {ct[0]['score']:.4f} "
        f"| {ct[1]['label_zh']} {ct[1]['score']:.4f} | {ct[2]['label_zh']} {ct[2]['score']:.4f} "
        f"| {c['latency_ms_internal_median']} | {g['latency_ms_internal_median']} |"
    )
L.append("")
L.append("## 3. Top-3、时延及异常处理测试记录")
L.append("")
L.append("### 3.1 时延汇总")
L.append("")
L.append("| 环境 | torch | 内部时延（预处理+H2D+前向+softmax+Top-3+D2H） | 墙钟时延 | 界面单次 |")
L.append("|---|---|---|---:|---:|")
L.append(f"| CPU | {cpu['torch']} | {rng(cpu['images'], 'latency_ms_internal_median')} ms "
         f"| {rng(cpu['images'], 'latency_ms_wall_median')} ms | 110.3 ms |")
L.append(f"| GPU | {gpu['torch']} | {rng(gpu['images'], 'latency_ms_internal_median')} ms "
         f"| {rng(gpu['images'], 'latency_ms_wall_median')} ms | 85.8 ms（首次） |")
L.append("")
L.append("GPU 首次界面推理包含 CUDA 预热，稳定后单张约 15–20 ms。")
L.append("")
L.append("### 3.2 异常处理")
L.append("")
L.append("| 用例 | 结果 |")
L.append("|---|---|")
L.append(f"| 损坏 JPG（文本改名） | {cpu['error_cases']['corrupt_image']} |")
L.append("| 未上传图片 | 提示“请先上传图片”，Top-3 为空 |")
L.append(f"| 缺失 checkpoint | state={cpu['error_cases']['missing_checkpoint_state']}；"
         f"识别按钮 disabled；{cpu['error_cases']['missing_checkpoint_predict']} |")
L.append("")
L.append("### 3.3 单元测试")
L.append("")
L.append("```")
L.append("Ran 60 tests in 19.265s")
L.append("OK (skipped=1)   # skip 项为 CUDA 专用计时用例")
L.append("```")
L.append("")
L.append("## 4. 发现的代码问题及必要修复")
L.append("")
L.append("| 问题 | 影响 | 处理 |")
L.append("|---|---|---|")
L.append("| 新增验收脚本文档字符串含 `\\.` 转义 | 仅产生 SyntaxWarning，不影响运行 | 改为 raw string，已修复 |")
L.append("| 项目默认 venv 安装的是 CPU 版 torch | 界面无法使用 GPU | 改用机器上已有的 cu124 环境完成 GPU 验收 |")
L.append("| GPU 首次推理 85.8 ms，稳定后约 20 ms | 首次点击偏慢 | CUDA 预热开销，非缺陷，已在文档注明 |")
L.append("")
L.append("integration/mvp-v1 的 checkpoint 严格加载、类别前缀校验、CUDA 同步计时均按预期工作，未发现阻断性问题。")
L.append("")
L.append("## 5. 最终结论")
L.append("")
L.append("可以用于课程项目展示：真实 EfficientNetV2-S 权重已加载并通过完整前向验证，"
         "8 张独立真实图片均返回 Top-3，异常路径不产生虚假预测，CPU 与 GPU 均可运行。")
L.append("")
L.append("对外表述必须遵守三条边界：")
L.append("")
L.append("1. 只能说“Demo 可运行、真实权重已接入”，不能说“识别准确率 100%”；")
L.append("2. 同源冻结测试集 1067/1067 不代表烤箱真实场景准确率；")
L.append("3. 未完成 ONNX 导出、RK3506 部署与真实烤箱外部测试，这些仍属后续工作。")
L.append("")
L.append("## 附录 A：八类冻结配置")
L.append("")
L.append("| ID | key | 中文名 |")
L.append("|---|---|---|")
for c in classes:
    L.append(f"| {c['id']} | `{c['key']}` | {c['zh']} |")
L.append("")
L.append("## 附录 B：复现命令")
L.append("")
L.append("```powershell")
L.append("cd D:\\天大港理工深圳学院\\AI&Project\\烤箱项目\\codex\\HeT_FoodAI_Demo")
L.append("$CKPT = (Get-ChildItem .\\runs\\food-v1.1.0-baseline -Filter best.pt -Recurse | Select-Object -First 1).FullName")
L.append(".\\.venv\\Scripts\\python.exe app\\app.py --check --checkpoint \"$CKPT\" --classes configs\\classes.json --device cpu")
L.append(".\\.venv\\Scripts\\python.exe app\\app.py --checkpoint \"$CKPT\" --classes configs\\classes.json --device cpu")
L.append("```")
L.append("")
L.append("## 附录 C：交付文件清单")
L.append("")
L.append("| 文件 | 说明 |")
L.append("|---|---|")
L.append("| `docs/ACCEPTANCE_REAL_MODEL_2026-10-21.md` | 验收记录（原始） |")
L.append("| `docs/acceptance-evidence/20261021/*.png` | 5 张界面截图 |")
L.append("| `docs/acceptance-evidence/20261021/acceptance.json` | CPU 机读结果 |")
L.append("| `docs/acceptance-evidence/20261021/acceptance_gpu.json` | GPU 机读结果 |")
L.append("| `scripts/acceptance_run.py` | 验收脚本 |")

md = OUT / "真实八类模型_Gradio_验收交付文档.md"
md.write_text("\n".join(L), encoding="utf-8")

# ------------------------------------------------------------------ DOCX
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Cm(2)
sec.left_margin = sec.right_margin = Cm(2.2)

style = doc.styles["Normal"]
style.font.name = "Microsoft YaHei"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

def h(text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
    return p

def para(text, bold=False, size=None, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    return p

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, x in enumerate(headers):
        hdr[i].text = ""
        r = hdr[i].paragraphs[0].add_run(x)
        r.bold = True
        r.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for i, x in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(x))
            r.font.size = Pt(9)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(TITLE)
r.bold = True
r.font.size = Pt(20)
r.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(SUBTITLE)
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
doc.add_paragraph()

para("本文档记录真实 EfficientNetV2-S 八类蔬菜模型接入 Gradio Demo 的端到端验收结果，"
     "按交付要求分为五项：运行截图、真实图片预测结果、Top-3 与时延及异常处理记录、"
     "代码问题与修复、课程展示结论。所有数字均为本机实测，softmax 分数不等同于模型准确率。")

h("0. 验收对象与环境", 1)
table(["项目", "内容"], [
    ["模型", f"{cpu['model_id']}（EfficientNetV2-S，stage 8）"],
    ["输入尺寸", f"{cpu['input_size']} px"],
    ["冻结类别配置", "configs/classes.json，status = frozen，version = v1.1.0"],
    ["class_list_sha256", cpu["class_list_sha256"]],
    ["checkpoint SHA256", cpu["checkpoint_sha256"]],
    ["CPU 环境", f"Windows · Python {cpu['python']} · torch {cpu['torch']} · device=cpu"],
    ["GPU 环境", f"Windows · Python {gpu['python']} · torch {gpu['torch']} · "
                 f"device=cuda（NVIDIA GeForce RTX 4050 Laptop GPU）"],
], widths=[4.2, 12.5])

h("1. Gradio 正常运行截图", 1)
for name, cap in [
    ("demo_ready_cucumber.png", "图 1  CPU 环境：真实黄瓜图片 Top-3，界面显示 110.3 ms"),
    ("demo_ready_gpu_tomato.png", "图 2  GPU 环境（device=cuda）：真实西红柿图片 Top-3"),
    ("demo_corrupt.png", "图 3  损坏图片：明确错误提示，无任何预测分数"),
    ("demo_no_upload.png", "图 4  未上传图片：提示“请先上传图片”"),
    ("demo_missing_checkpoint.png", "图 5  未配置模型：显示未就绪，识别按钮禁用"),
]:
    path = EVID / name
    if path.exists():
        doc.add_picture(str(path), width=Cm(15.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        para(cap, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

h("2. 真实图片预测结果", 1)
para("8 张独立真实图片，CPU 与 GPU 预测结果一致；每张重复 5 次（CPU）/ 10 次（GPU），"
     "时延为中位数，单位 ms。")
rows = []
for f, c in cpu_by.items():
    g = gpu_by[f]
    ct = c["top3"]
    rows.append([f, f"{ct[0]['label_zh']} {ct[0]['score']:.4f}",
                 f"{ct[1]['label_zh']} {ct[1]['score']:.4f}",
                 f"{ct[2]['label_zh']} {ct[2]['score']:.4f}",
                 c["latency_ms_internal_median"], g["latency_ms_internal_median"]])
table(["图片", "Top-1", "Top-2", "Top-3", "CPU ms", "GPU ms"], rows,
      widths=[4.6, 3.2, 3.0, 2.8, 1.6, 1.6])

h("3. Top-3、时延及异常处理测试记录", 1)
h("3.1 时延汇总", 2)
table(["环境", "torch", "内部时延", "墙钟时延", "界面单次"], [
    ["CPU", cpu["torch"], f"{rng(cpu['images'], 'latency_ms_internal_median')} ms",
     f"{rng(cpu['images'], 'latency_ms_wall_median')} ms", "110.3 ms"],
    ["GPU", gpu["torch"], f"{rng(gpu['images'], 'latency_ms_internal_median')} ms",
     f"{rng(gpu['images'], 'latency_ms_wall_median')} ms", "85.8 ms（首次）"],
], widths=[1.4, 3.4, 4.2, 4.2, 3.5])
para("GPU 首次界面推理包含 CUDA 预热，稳定后单张约 15–20 ms。")

h("3.2 异常处理", 2)
table(["用例", "结果"], [
    ["损坏 JPG（文本改名）", cpu["error_cases"]["corrupt_image"]],
    ["未上传图片", "提示“请先上传图片”，Top-3 为空"],
    ["缺失 checkpoint", f"state={cpu['error_cases']['missing_checkpoint_state']}；"
                        f"识别按钮 disabled；{cpu['error_cases']['missing_checkpoint_predict']}"],
], widths=[4.5, 12.2])

h("3.3 单元测试", 2)
para("Ran 60 tests in 19.265s；OK (skipped=1)，skip 项为 CUDA 专用计时用例。")

h("4. 发现的代码问题及必要修复", 1)
table(["问题", "影响", "处理"], [
    ["新增验收脚本文档字符串含 \\. 转义", "仅产生 SyntaxWarning，不影响运行", "改为 raw string，已修复"],
    ["项目默认 venv 安装的是 CPU 版 torch", "界面无法使用 GPU", "改用机器上已有的 cu124 环境完成 GPU 验收"],
    ["GPU 首次推理 85.8 ms，稳定后约 20 ms", "首次点击偏慢", "CUDA 预热开销，非缺陷，已在文档注明"],
], widths=[6.0, 5.2, 5.5])
para("integration/mvp-v1 的 checkpoint 严格加载、类别前缀校验、CUDA 同步计时均按预期工作，"
     "未发现阻断性问题。")

h("5. 最终结论", 1)
para("可以用于课程项目展示：真实 EfficientNetV2-S 权重已加载并通过完整前向验证，"
     "8 张独立真实图片均返回 Top-3，异常路径不产生虚假预测，CPU 与 GPU 均可运行。")
para("对外表述必须遵守三条边界：", bold=True)
for text in [
    "1. 只能说“Demo 可运行、真实权重已接入”，不能说“识别准确率 100%”；",
    "2. 同源冻结测试集 1067/1067 不代表烤箱真实场景准确率；",
    "3. 未完成 ONNX 导出、RK3506 部署与真实烤箱外部测试，这些仍属后续工作。",
]:
    para(text)

h("附录 A：八类冻结配置", 1)
table(["ID", "key", "中文名"], [[c["id"], c["key"], c["zh"]] for c in classes],
      widths=[1.8, 5.0, 4.0])

h("附录 B：复现命令", 1)
for line in [
    "cd D:\\天大港理工深圳学院\\AI&Project\\烤箱项目\\codex\\HeT_FoodAI_Demo",
    "$CKPT = (Get-ChildItem .\\runs\\food-v1.1.0-baseline -Filter best.pt -Recurse | Select-Object -First 1).FullName",
    ".\\.venv\\Scripts\\python.exe app\\app.py --check --checkpoint \"$CKPT\" --classes configs\\classes.json --device cpu",
    ".\\.venv\\Scripts\\python.exe app\\app.py --checkpoint \"$CKPT\" --classes configs\\classes.json --device cpu",
]:
    p = doc.add_paragraph()
    r = p.add_run(line)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)

h("附录 C：交付文件清单", 1)
table(["文件", "说明"], [
    ["docs/ACCEPTANCE_REAL_MODEL_2026-10-21.md", "验收记录（原始）"],
    ["docs/acceptance-evidence/20261021/*.png", "5 张界面截图"],
    ["docs/acceptance-evidence/20261021/acceptance.json", "CPU 机读结果"],
    ["docs/acceptance-evidence/20261021/acceptance_gpu.json", "GPU 机读结果"],
    ["scripts/acceptance_run.py", "验收脚本"],
], widths=[8.5, 8.0])

docx_path = OUT / "真实八类模型_Gradio_验收交付文档.docx"
doc.save(docx_path)
print(md)
print(docx_path)

